"""Static layout check of the PPTX: wraps every text box with real Inter metrics and reports boxes whose text needs more height
than the box has, text smaller than 14 pt, and shapes that leave the slide.   python check_fit.py <deck.pptx> <inter-dir>"""
import sys
from pptx import Presentation
from pptx.util import Emu
from PIL import ImageFont

deck, fontdir = sys.argv[1], sys.argv[2]
FONTS = {}


def font(size, bold):
    key = (round(size * 4), bold)
    if key not in FONTS:
        FONTS[key] = ImageFont.truetype(f'{fontdir}/Inter-{"Bold" if bold else "Regular"}.ttf', size=size * 4)
    return FONTS[key]


def wrap_height(tf, width_pt):
    total = 0.0
    for p in tf.paragraphs:
        runs = [(r.text, r.font.size.pt if r.font.size else 18, bool(r.font.bold)) for r in p.runs]
        if not runs: continue
        size = max(r[1] for r in runs); ls = p.line_spacing if isinstance(p.line_spacing, float) else 1.0
        words = []
        for t, sz, b in runs:
            for w in t.replace('\n', ' ').split(' '):
                words.append((w, sz, b))
        lines, cur = 1, 0.0
        for w, sz, b in words:
            wlen = font(sz, b).getlength(w + ' ') / 4
            if cur + wlen - font(sz, b).getlength(' ') / 4 > width_pt and cur > 0:
                lines += 1; cur = 0.0
            cur += wlen
        total += lines * size * 1.21 * ls + (p.space_after.pt if p.space_after else 0)
    return total


prs = Presentation(deck); problems = 0
for n, slide in enumerate(prs.slides, 1):
    for sh in slide.shapes:
        if sh.left < 0 or sh.top < 0 or sh.left + sh.width > prs.slide_width + 10 or sh.top + sh.height > prs.slide_height + 10:
            print(f'slide {n}: shape {sh.shape_type} {sh.name} outside the slide'); problems += 1
        if sh.has_text_frame and sh.text_frame.text.strip():
            w = Emu(sh.width).pt; h = Emu(sh.height).pt
            need = wrap_height(sh.text_frame, w)
            sizes = [r.font.size.pt for p in sh.text_frame.paragraphs for r in p.runs if r.font.size]
            if sizes and min(sizes) < 14: print(f'slide {n}: {min(sizes)} pt text "{sh.text_frame.text[:40]}"'); problems += 1
            if need > h + 2:
                print(f'slide {n}: text needs {need:.0f} pt but box is {h:.0f} pt: "{sh.text_frame.text[:50]}"'); problems += 1
print('problems:', problems)
