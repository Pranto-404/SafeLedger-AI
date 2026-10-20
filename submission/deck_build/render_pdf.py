"""Renders the deck PPTX to a 16:9 PDF with embedded fonts and working links, without Office or LibreOffice.

It walks the PPTX that build_deck.py wrote (rectangles, rounded rectangles, chevrons, lines, text boxes, pictures) and lays the same
geometry out as HTML, then prints it with Chrome. Only the shape types used by this deck are supported.
    python render_pdf.py deck.pptx out.pdf
"""
import html, sys, tempfile
from pathlib import Path
from pptx import Presentation
from pptx.util import Emu
from pptx.enum.shapes import MSO_SHAPE_TYPE
from playwright.sync_api import sync_playwright
from pptx.oxml.ns import qn

deck, out = sys.argv[1], sys.argv[2]
prs = Presentation(deck)
LINE_H = 1.21  # Inter's natural line height (ascender + descender), as PowerPoint uses it


def pt(v): return Emu(v).pt


def color(fmt):
    try:
        return '#' + str(fmt.rgb)
    except Exception:
        return None


def shape_css(sh):
    x, y, w, h = pt(sh.left), pt(sh.top), pt(sh.width), pt(sh.height)
    return f'left:{x:.2f}pt;top:{y:.2f}pt;width:{w:.2f}pt;height:{h:.2f}pt;'


def fill_line_css(sh):
    css = ''
    f = sh.fill
    if f.type == 1: css += f'background:{color(f.fore_color)};'
    ln = sh.line
    if ln.fill.type == 1: css += f'border:{ln.width.pt if ln.width else 1}pt solid {color(ln.color)};box-sizing:border-box;'
    return css


def render_text(sh):
    tf = sh.text_frame
    anchor = {1: 'flex-start', 3: 'center', 4: 'flex-end'}.get(int(tf.vertical_anchor) if tf.vertical_anchor else 1, 'flex-start')
    paras = []
    for p in tf.paragraphs:
        ls = p.line_spacing if isinstance(p.line_spacing, float) else 1.0
        al = {1: 'left', 2: 'center', 3: 'right'}.get(int(p.alignment) if p.alignment else 1, 'left')
        sa = p.space_after.pt if p.space_after else 0
        runs = []
        for r in p.runs:
            f = r.font; rpr = r._r.rPr
            spc = (int(rpr.get('spc')) / 100) if rpr is not None and rpr.get('spc') else 0
            st = f'font-size:{f.size.pt}pt;font-weight:{700 if f.bold else 400};color:{color(f.color)};'
            if spc: st += f'letter-spacing:{spc}pt;'
            t = html.escape(r.text).replace('\n', '<br>')
            link = r.hyperlink.address
            runs.append(f'<a href="{html.escape(link)}" style="{st}text-decoration:none">{t}</a>' if link else f'<span style="{st}">{t}</span>')
        mx = max([r.font.size.pt for r in p.runs] or [18])
        paras.append(f'<p style="margin:0 0 {sa}pt 0;text-align:{al};line-height:{mx * LINE_H * ls:.2f}pt;font-size:{mx}pt">{"".join(runs) or "&nbsp;"}</p>')
    return f'<div style="position:absolute;{shape_css(sh)}display:flex;flex-direction:column;justify-content:{anchor};overflow:visible">{"".join(paras)}</div>'


def render_shape(sh, tmp):
    t = sh.shape_type
    if t == MSO_SHAPE_TYPE.PICTURE:
        p = Path(tmp) / f'img{sh.shape_id}_{id(sh)}.{sh.image.ext}'; p.write_bytes(sh.image.blob)
        img = f'<img src="{p.as_uri()}" style="width:100%;height:100%;display:block">'
        link = sh.click_action.hyperlink.address
        inner = f'<a href="{html.escape(link)}">{img}</a>' if link else img
        return f'<div style="position:absolute;{shape_css(sh)}">{inner}</div>'
    if t == MSO_SHAPE_TYPE.TEXT_BOX:
        return render_text(sh)
    if t == MSO_SHAPE_TYPE.LINE:
        x1, y1, x2, y2 = pt(sh.begin_x), pt(sh.begin_y), pt(sh.end_x), pt(sh.end_y)
        w = sh.line.width.pt if sh.line.width else 1
        return f'<div style="position:absolute;left:{x1:.2f}pt;top:{y1:.2f}pt;width:{x2 - x1:.2f}pt;height:{w}pt;background:{color(sh.line.color)}"></div>'
    if t == MSO_SHAPE_TYPE.AUTO_SHAPE:
        name = sh.auto_shape_type.name if sh.auto_shape_type else ''
        css = shape_css(sh) + fill_line_css(sh)
        if 'ROUNDED' in name:
            r = sh.adjustments[0] * min(pt(sh.width), pt(sh.height))
            css += f'border-radius:{r:.2f}pt;'
        elif name == 'CHEVRON':
            d = 0.5 * min(pt(sh.width), pt(sh.height))
            css += f'clip-path:polygon(0 0,calc(100% - {d:.2f}pt) 0,100% 50%,calc(100% - {d:.2f}pt) 100%,0 100%,{d:.2f}pt 50%);'
            if sh.rotation: css += f'transform:rotate({sh.rotation}deg);'
        body = render_text(sh) if sh.has_text_frame and sh.text_frame.text.strip() else ''
        return f'<div style="position:absolute;{css}"></div>{body}'
    raise SystemExit(f'unsupported shape {t} {sh.name}')


with tempfile.TemporaryDirectory() as tmp:
    pages = []
    for slide in prs.slides:
        bg = '#' + str(slide.background.fill.fore_color.rgb)
        body = ''.join(render_shape(sh, tmp) for sh in slide.shapes)
        pages.append(f'<section class="slide" style="background:{bg}">{body}</section>')
    fonts = Path(__file__).resolve().parent / 'fonts'
    faces = ''.join(f"@font-face{{font-family:'Inter';font-weight:{w};src:url('{(fonts / f).as_uri()}')}}" for w, f in [(400, 'Inter-Regular.ttf'), (700, 'Inter-Bold.ttf')])
    doc = ('<!doctype html><html><head><meta charset="utf-8"><title>Capzoryx Copilot: 10-minute competition pitch</title><style>' + faces +
           '@page{size:960pt 540pt;margin:0}html,body{margin:0;padding:0;background:#000}'
           '.slide{position:relative;width:960pt;height:540pt;overflow:hidden;page-break-after:always;font-family:Inter,sans-serif}'
           '.slide *{box-sizing:border-box;font-family:Inter,sans-serif}</style></head><body>' + ''.join(pages) + '</body></html>')
    src = Path(tmp) / 'deck.html'; src.write_text(doc, encoding='utf-8')
    with sync_playwright() as pw:
        b = pw.chromium.launch(channel='chrome'); pg = b.new_page()
        pg.goto(src.as_uri()); pg.wait_for_timeout(1500); pg.evaluate('document.fonts.ready')
        pg.pdf(path=out, width='13.3333in', height='7.5in', print_background=True, prefer_css_page_size=True)
        b.close()
import fitz  # document properties (Chrome leaves author and subject empty)
doc = fitz.open(out)
doc.set_metadata({'title': 'Capzoryx Copilot: 10-minute competition pitch', 'author': 'Team Capzoryx', 'subject': '2026 Shenzhen International FinTechathon, International Track, Topic E: SME Finance Copilot', 'creator': 'build_deck.py + render_pdf.py', 'producer': 'Chrome (print to PDF) / PyMuPDF'})
tmp_out = out + '.tmp'; doc.save(tmp_out, garbage=3, deflate=True); doc.close(); Path(tmp_out).replace(out)
print('wrote', out)
